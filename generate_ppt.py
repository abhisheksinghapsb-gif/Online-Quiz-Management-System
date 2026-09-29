import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    NAVY = RGBColor(15, 23, 42)       # #0f172a
    DARK_BLUE = RGBColor(30, 41, 59)   # #1e293b
    ROYAL_BLUE = RGBColor(37, 99, 235) # #2563eb
    LIGHT_BG = RGBColor(248, 250, 252) # #f8fafc
    WHITE = RGBColor(255, 255, 255)
    TEXT_DARK = RGBColor(30, 41, 59)
    TEXT_MUTED = RGBColor(100, 116, 139)
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240)
    SUCCESS_GREEN = RGBColor(16, 185, 129)
    ACCENT_PURPLE = RGBColor(139, 92, 246)
    AMBER = RGBColor(245, 158, 11)

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="CIE-2: SKILL DEVELOPMENT LABORATORY USING JAVA (BIT25434A0X)"):
        # Top accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ROYAL_BLUE
        bar.line.fill.background()

        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = "Arial"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = ROYAL_BLUE
        p0.space_after = Pt(4)

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = "Arial"
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, title, items, badge_text=None, border_color=CARD_BORDER, bg_color=CARD_BG):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.25)

        if badge_text:
            p_badge = tf.paragraphs[0]
            p_badge.text = badge_text.upper()
            p_badge.font.name = "Arial"
            p_badge.font.size = Pt(9)
            p_badge.font.bold = True
            p_badge.font.color.rgb = ROYAL_BLUE
            p_badge.space_after = Pt(2)
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.name = "Arial"
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        p_title.space_after = Pt(10)

        for itm in items:
            p = tf.add_paragraph()
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.space_after = Pt(6)
            
            if isinstance(itm, tuple):
                bold_prefix, rest = itm
                run_b = p.add_run()
                run_b.text = "• " + bold_prefix + ": "
                run_b.font.bold = True
                run_b.font.color.rgb = NAVY
                run_r = p.add_run()
                run_r.text = rest
                run_r.font.color.rgb = TEXT_DARK
            else:
                p.text = "• " + itm
                p.font.color.rgb = TEXT_DARK

        return shape

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Blue Modern Layout)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, NAVY)

    # Accent glow top bar
    bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = ROYAL_BLUE
    bar1.line.fill.background()

    # Title & College info
    tb_title = slide1.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(3.0))
    tf1 = tb_title.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "ARMY INSTITUTE OF TECHNOLOGY, PUNE"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.space_after = Pt(6)

    p = tf1.add_paragraph()
    p.text = "DEPARTMENT OF INFORMATION TECHNOLOGY"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(96, 165, 250)
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "Online Quiz Management System"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(10)

    p = tf1.add_paragraph()
    p.text = "CIE–2: Practical Problem Solving in Java (Unit III & Unit IV) | Class: SE IT B"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(203, 213, 225)

    # Cards for Examiner and Team Members
    # Left Card: Examiner / In-charge
    card_guide = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(4.0), Inches(2.6))
    card_guide.fill.solid()
    card_guide.fill.fore_color.rgb = DARK_BLUE
    card_guide.line.color.rgb = RGBColor(51, 65, 85)
    tfg = card_guide.text_frame
    tfg.margin_left = tfg.margin_right = Inches(0.3)
    tfg.margin_top = Inches(0.3)
    pg = tfg.paragraphs[0]
    pg.text = "UNDER THE GUIDANCE OF"
    pg.font.size = Pt(11)
    pg.font.bold = True
    pg.font.color.rgb = RGBColor(148, 163, 184)
    pg.space_after = Pt(8)
    pg2 = tfg.add_paragraph()
    pg2.text = "Mrs. Trupti Najan"
    pg2.font.size = Pt(18)
    pg2.font.bold = True
    pg2.font.color.rgb = WHITE
    pg2.space_after = Pt(4)
    pg3 = tfg.add_paragraph()
    pg3.text = "Assistant Professor\nDepartment of Information Technology\nArmy Institute of Technology, Pune"
    pg3.font.size = Pt(11)
    pg3.font.color.rgb = RGBColor(203, 213, 225)

    # Right Card: Group Members (IT B)
    card_team = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.1), Inches(4.2), Inches(7.4), Inches(2.6))
    card_team.fill.solid()
    card_team.fill.fore_color.rgb = DARK_BLUE
    card_team.line.color.rgb = RGBColor(51, 65, 85)
    tft = card_team.text_frame
    tft.margin_left = tft.margin_right = Inches(0.3)
    tft.margin_top = Inches(0.3)
    pt = tft.paragraphs[0]
    pt.text = "PROJECT TEAM MEMBERS (CLASS: IT B)"
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = RGBColor(148, 163, 184)
    pt.space_after = Pt(10)

    members = [
        ("Aditya Yadav", "Roll No: 8108", "System Architecture, Abstract Models & Aggregation"),
        ("Abhishekh Singh", "Roll No: 8104", "Polymorphic Question Archetypes & Dynamic Binding"),
        ("Priyam Raj", "Roll No: 8134", "Interfaces & Pluggable Scoring Strategy Policies"),
        ("Utkarsh Chauhan", "Roll No: 8154", "Custom Checked Exception Hierarchy & Defensive Validation")
    ]
    for name, roll, role in members:
        p_m = tft.add_paragraph()
        p_m.space_after = Pt(4)
        r1 = p_m.add_run()
        r1.text = f"• {name} ({roll}) — "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = WHITE
        r2 = p_m.add_run()
        r2.text = role
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # SLIDE 2: Problem Statement & Motivation
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, LIGHT_BG)
    add_header(slide2, "Problem Statement & Project Motivation")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.9),
             "Formal Problem Statement (As required by CIE-2 Guidelines)",
             [
                 "Traditional manual examinations and fragmented testing tools suffer from subjective grading inaccuracies, administrative overhead, and brittle input parsing that crashes during runtime errors.",
                 "The Online Quiz Management System is an end-to-end, object-oriented assessment engine developed in Core Java. It enables faculty to design, categorize, and administer heterogeneous quizzes (MCQ, True/False, and Numeric questions) while providing students with automated evaluation, dynamic letter grading, and comprehensive performance scorecards under interchangeable scoring models.",
                 "The solution enforces complete domain integrity through custom checked exceptions and polymorphic dispatch, completely eliminating unhandled runtime crashes."
             ], badge_text="1 Short Paragraph Problem Formulation", bg_color=WHITE)

    add_card(slide2, Inches(0.8), Inches(3.8), Inches(5.7), Inches(3.2),
             "Academic & Practical Challenges",
             [
                 ("Subjective Scoring Inconsistencies", "Difficulty in standardizing linear grading vs competitive negative marking."),
                 ("Rigid Question Paradigms", "Standard systems hardcode single question types rather than supporting polymorphic question sets."),
                 ("Brittle Console Execution", "User typographic errors (e.g. entering non-numeric characters) typically terminate standard programs.")
             ], badge_text="The Challenge")

    add_card(slide2, Inches(6.8), Inches(3.8), Inches(5.7), Inches(3.2),
             "Core Java Problem-Solving Approach",
             [
                 ("Meaningful OOP Application", "Uses Abstract Classes for question models and Interfaces for evaluation strategies."),
                 ("Strict Defensive Design", "A dedicated 5-tier exception hierarchy isolates illegal bounds, options, and duplicate IDs."),
                 ("Dual Mode Delivery", "Operates as both an interactive Terminal CLI and a modern Browser Application via embedded Java server.")
             ], badge_text="The Solution", border_color=ROYAL_BLUE)

    # =========================================================================
    # SLIDE 3: Objectives & Scope
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, LIGHT_BG)
    add_header(slide3, "Project Objectives & Technical Scope")

    cards_s3 = [
        ("1. Dynamic Question Archetypes",
         [("Polymorphic Model", "Implement an abstract base class Question subclassed into MultipleChoiceQuestion, TrueFalseQuestion, and NumericQuestion."),
          ("Encapsulated State", "Uniformly manage IDs, marks, syllabus topics, and difficulty levels across all question types.")],
         Inches(0.8), Inches(1.6)),

        ("2. Decoupled Service Contracts",
         [("Interface Contracts", "Define QuizOperations and QuizEvaluator interfaces to isolate presentation from core business logic."),
          ("Strategy Design Pattern", "Permit runtime switching between Standard Linear Grading and Competitive Negative Marking.")],
         Inches(6.8), Inches(1.6)),

        ("3. Comprehensive Fault Tolerance",
         [("Custom Exception Tree", "Root all errors under QuizException (QuizNotFoundException, DuplicateQuizException, InvalidOptionException, etc.)."),
          ("Zero-Crash Guarantee", "Intercept typographic mistakes and out-of-bound inputs, prompting user recovery without session loss.")],
         Inches(0.8), Inches(4.4)),

        ("4. Real-Time Feedback & Analytics",
         [("Detailed Scorecard", "Generate instant audit reports showing earned points, correct answers, net percentages, and letter grades."),
          ("Continuous Project Scope", "Satisfies CIE-2 Phase 1 prototype and lays foundation for Phase 2/3 persistence and GUI.")],
         Inches(6.8), Inches(4.4))
    ]

    for title, items, left, top in cards_s3:
        add_card(slide3, left, top, Inches(5.7), Inches(2.6), title, items, badge_text="Objective")

    # =========================================================================
    # SLIDE 4: System Architecture & Package Structure
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, LIGHT_BG)
    add_header(slide4, "System Architecture & Package Hierarchy")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.4),
             "Package Organization",
             [
                 ("com.ait.quiz.model", "Question (Abstract), MultipleChoiceQuestion, TrueFalseQuestion, NumericQuestion, User (Abstract), Student, Instructor, Quiz, QuizAttempt, DifficultyLevel"),
                 ("com.ait.quiz.service", "QuizOperations (Interface), QuizEvaluator (Interface), StandardGradingPolicy, NegativeMarkingGradingPolicy, QuizManager"),
                 ("com.ait.quiz.exception", "QuizException (Base Checked), QuizNotFoundException, DuplicateQuizException, InvalidOptionException, InvalidQuestionException, EmptyQuizException"),
                 ("com.ait.quiz.util", "InputValidator (Defensive scanner), ConsoleUI (Formatting)"),
                 ("com.ait.quiz.main", "QuizApplication (Menu & test harness)"),
                 ("com.ait.quiz.web", "QuizWebServer (Embedded Java HTTP server)")
             ], badge_text="Modular Package Design")

    add_card(slide4, Inches(5.6), Inches(1.6), Inches(6.9), Inches(2.55),
             "Multi-Tier Architectural Separation",
             [
                 ("Presentation Layer", "Interactive CLI Console (`QuizApplication`) and Modern Browser UI (`index.html`) communicating over JSON REST endpoints."),
                 ("Business Service Layer", "Encapsulates repository operations via `QuizOperations` and evaluation algorithms via `QuizEvaluator`."),
                 ("Domain Model Layer", "Encapsulates state, integrity constraints, and dynamic question validation logic.")
             ], badge_text="Layered Architecture")

    add_card(slide4, Inches(5.6), Inches(4.45), Inches(6.9), Inches(2.55),
             "Core Design Principles Applied",
             [
                 ("Separation of Concerns", "Presentation code never manipulates raw question arrays directly; all workflows pass through interface contracts."),
                 ("Open / Closed Principle", "System is open for extension (e.g. adding FillInTheBlankQuestion or relative grading) without modifying existing code."),
                 ("Liskov Substitution", "Any specialized Question subclass substitutes seamlessly inside a Quiz collection.")
             ], badge_text="Engineering Rigor")

    # =========================================================================
    # SLIDE 5: Unit III Concept 1 - Abstract Classes
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, LIGHT_BG)
    add_header(slide5, "Java Unit III: Abstract Classes (Question & User)")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Abstract Base Class: Question",
             [
                 ("Why an Abstract Class?", "All questions share identical state (id, questionText, marks, topic, difficulty) but a generic question cannot be displayed or evaluated without specific choice logic."),
                 ("Encapsulated State", "Private fields with public getters/setters and defensive validation constructor."),
                 ("Concrete Common Method", "displayHeader() prints formatted ID, marks badge, topic, and difficulty tag uniformly."),
                 ("Abstract Method 1", "abstract void displayQuestion(); — forces subclasses to render custom choices."),
                 ("Abstract Method 2", "abstract boolean checkAnswer(String ans) throws InvalidOptionException; — executes subclass-specific answer parsing."),
                 ("Abstract Method 3", "abstract String getCorrectAnswerFormatted(); — returns human-readable solution."),
                 ("Instantiation Guard", "Prevents calling 'new Question(...)'; only concrete subclasses can be instantiated.")
             ], badge_text="Core Domain Abstraction")

    add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Concrete Subclasses & User Hierarchy",
             [
                 ("MultipleChoiceQuestion", "Extends Question. Stores List<String> options and correctOptionIndex. Overrides checkAnswer() to parse single-character options (A–D)."),
                 ("TrueFalseQuestion", "Extends Question. Stores boolean correctAnswer. Overrides checkAnswer() accepting T/F/True/False input."),
                 ("NumericQuestion", "Extends Question. Stores double correctAnswer and tolerance delta for mathematical floating-point comparisons."),
                 ("Abstract Class: User", "Encapsulates common user identity (userId, name, email). Declares abstract void displayDashboard()."),
                 ("Student (Subclass)", "Overrides displayDashboard() to render completed quiz count, average score percentage, and attempt history."),
                 ("Instructor (Subclass)", "Overrides displayDashboard() displaying faculty designation, department, and authoring capabilities.")
             ], badge_text="Inheritance & Specialization")

    # =========================================================================
    # SLIDE 6: Unit III Concept 2 - Interfaces & Strategy Pattern
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, LIGHT_BG)
    add_header(slide6, "Java Unit III: Interfaces & Pluggable Strategy Pattern")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Interface: QuizOperations",
             [
                 ("Interface Purpose", "Defines the complete behavioral contract for managing quizzes and tracking participant attempts without tying to storage."),
                 ("Lifecycle Contracts", "createQuiz(), addQuestionToQuiz(), getQuiz(), deleteQuiz()."),
                 ("Exception Declarations", "Declares checked exceptions in signatures (throws QuizNotFoundException, DuplicateQuizException)."),
                 ("Service Realization", "Implemented by QuizManager using Collections Framework (LinkedHashMap<String, Quiz> and List<QuizAttempt>)."),
                 ("Future Scalability", "Enables writing a DatabaseQuizManager implementing QuizOperations in Phase 3 without altering the UI.")
             ], badge_text="CRUD Lifecycle Contract")

    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Interface: QuizEvaluator (Strategy Pattern)",
             [
                 ("The Strategy Design Pattern", "Decouples the test execution engine from scoring algorithms; grading policy is injected at runtime."),
                 ("Contract Methods", "evaluateScore(QuizAttempt), evaluateScore(marks, wrong, penalty), generateGrade(percentage), printDetailedReport()."),
                 ("StandardGradingPolicy", "Implements QuizEvaluator. Full marks for correct answers, zero deduction for incorrect answers (University Linear Model)."),
                 ("NegativeMarkingGradingPolicy", "Implements QuizEvaluator. Penalizes wrong answers with configurable penalty rate (e.g. 25% negative marking for competitive exams)."),
                 ("Rubric Alignment", "Directly fulfills the syllabus requirement for pluggable grading behavior.")
             ], badge_text="Interchangeable Grading Strategies")

    # =========================================================================
    # SLIDE 7: Unit III Concept 3 - Polymorphism in Action
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, LIGHT_BG)
    add_header(slide7, "Java Unit III: Polymorphism (Runtime & Compile-Time)")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Runtime Polymorphism (Dynamic Method Dispatch)",
             [
                 ("Polymorphic Question Iteration", "In QuizApplication, questions are stored as List<Question>. When conducting a quiz:"),
                 ("Method Overriding", "q.displayQuestion() and q.checkAnswer(ans) dynamically bind at runtime to MCQ, TrueFalse, or Numeric logic without any 'instanceof' or switch checks."),
                 ("Interface Polymorphism", "QuizEvaluator evaluator = (mode == 1) ? new StandardGradingPolicy() : new NegativeMarkingGradingPolicy(); evaluator.evaluateScore(attempt);"),
                 ("Polymorphic Dashboard", "User userRef = currentStudent; userRef.displayDashboard(); dynamically resolves to Student's dashboard implementation."),
                 ("JVM Virtual Method Table", "Calls resolve dynamically via JVM vtable based on the actual object on heap.")
             ], badge_text="Dynamic Binding")

    add_card(slide7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Compile-Time Polymorphism (Method Overloading)",
             [
                 ("Method Overloading in QuizOperations", "Multiple search signatures provide flexible querying options for students:"),
                 ("Overloaded Search 1", "List<Quiz> searchQuiz(String topic); — filters quizzes matching a topic keyword."),
                 ("Overloaded Search 2", "List<Quiz> searchQuiz(String topic, DifficultyLevel level); — filters by topic and difficulty level simultaneously."),
                 ("Method Overloading in InputValidator", "readInteger(Scanner sc, String prompt); and readInteger(Scanner sc, String prompt, int min, int max);"),
                 ("Method Overloading in QuizEvaluator", "evaluateScore(QuizAttempt attempt); vs evaluateScore(double earnedMarks, int wrongCount, double penalty);")
             ], badge_text="Static Binding / Overloading")

    # =========================================================================
    # SLIDE 8: Unit IV Concept 4 - Exception Handling Hierarchy
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, LIGHT_BG)
    add_header(slide8, "Java Unit IV: Robust Exception Handling Hierarchy")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Custom Checked Exception Hierarchy",
             [
                 ("Base Class: QuizException", "Extends java.lang.Exception. Root checked exception for domain integrity."),
                 ("QuizNotFoundException", "Thrown when accessing or deleting an unmapped Quiz ID (e.g. 'GHOST-404')."),
                 ("DuplicateQuizException", "Thrown when faculty attempts to register a quiz with an existing ID (e.g. 'JAVA-OOP')."),
                 ("InvalidQuestionException", "Thrown on constraint failure: blank prompt, marks <= 0, or MCQ options < 2."),
                 ("InvalidOptionException", "Thrown when a student provides an out-of-range option (e.g. 'Z' on MCQ or text in numeric)."),
                 ("EmptyQuizException", "Thrown if an attempt is made to conduct a quiz containing zero questions.")
             ], badge_text="Custom Exception Tree")

    add_card(slide8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Structured Try-Catch-Finally Architecture",
             [
                 ("The 'try' Block", "Isolates critical operations: reading console inputs, invoking domain services, and validating student answers."),
                 ("The 'catch' Block", "Intercepts specific checked exceptions, prints clear diagnostic recovery guidance, and re-prompts without terminating."),
                 ("The 'finally' Block", "Guaranteed execution regardless of exception or normal completion: manages graceful resource shutdown and audit completion logs."),
                 ("Throw & Throws Keywords", "Methods declare throws in signatures; business rules explicitly throw exceptions on boundary violations."),
                 ("Defensive Input Handling", "InputValidator wraps NumberFormatException, guaranteeing zero console crashes.")
             ], badge_text="Fault-Tolerant Execution")

    # =========================================================================
    # SLIDE 9: UML Class & Workflow Diagrams
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, LIGHT_BG)
    add_header(slide9, "UML Class Relationships & Execution Sequence")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "UML Structural Class Diagram",
             [
                 ("Inheritance (Generalization)", "Question <|-- MultipleChoiceQuestion\nQuestion <|-- TrueFalseQuestion\nQuestion <|-- NumericQuestion\nUser <|-- Student, User <|-- Instructor"),
                 ("Interface Realization", "QuizOperations <|.. QuizManager\nQuizEvaluator <|.. StandardGradingPolicy\nQuizEvaluator <|.. NegativeMarkingGradingPolicy"),
                 ("Aggregation & Association", "Quiz '1' o-- '*' Question (A quiz aggregates polymorphic questions)\nQuizManager '1' *-- '*' Quiz\nStudent '1' *-- '*' QuizAttempt"),
                 ("Exception Dependency", "QuizOperations and Question methods declare dependency on QuizException hierarchy.")
             ], badge_text="OOP Relationships")

    add_card(slide9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.4),
             "Runtime Conduction Sequence Flow",
             [
                 ("1. Launch & Selection", "Student chooses quiz ID -> QuizManager validates existence (throws QuizNotFoundException if absent)."),
                 ("2. Readiness Validation", "Quiz verifies questionCount > 0 (throws EmptyQuizException if blank)."),
                 ("3. Strategy Binding", "Student selects scoring policy -> QuizEvaluator instance created."),
                 ("4. Polymorphic Iteration", "Loop over List<Question> -> q.displayQuestion() -> student inputs choice -> q.checkAnswer()."),
                 ("5. Exception Recovery", "If student enters illegal format -> catches InvalidOptionException and re-prompts seamlessly."),
                 ("6. Report Generation", "QuizEvaluator calculates net score -> records QuizAttempt -> prints itemized scorecard.")
             ], badge_text="Execution Lifecycle")

    # =========================================================================
    # SLIDE 10: Implementation & Live Demonstration
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, LIGHT_BG)
    add_header(slide10, "Dual Mode Implementation & Live Demonstration")

    add_card(slide10, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.4),
             "Mode 1: Interactive Console CLI",
             [
                 ("Launcher Script", "run.bat"),
                 ("Pure Core Java", "Works directly in Windows Command Prompt without external dependencies."),
                 ("ASCII Boxed Menus", "Professional borders, separated student/faculty dashboards."),
                 ("Graceful Recovery", "Try entering 'Z' on MCQ question 1 — system alerts and re-prompts."),
                 ("Detailed Audit", "Generates full item-by-item breakdown with earned marks and letter grades.")
             ], badge_text="Terminal Interface")

    add_card(slide10, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.4),
             "Mode 2: Modern Web Application",
             [
                 ("Launcher Script", "run_web.bat"),
                 ("Embedded Server", "Custom QuizWebServer built using standard com.sun.net.httpserver.HttpServer."),
                 ("Zero External JARs", "Runs out-of-the-box on standard JDK on http://localhost:8080."),
                 ("Interactive UI", "10-minute live countdown timer, radio options, instant scorecard modal."),
                 ("Faculty Panel", "Dynamic modals to create quizzes, add questions, and view submissions.")
             ], badge_text="Browser Web Interface", border_color=ROYAL_BLUE)

    add_card(slide10, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.4),
             "Mode 3: Automated Test Matrix",
             [
                 ("Launcher Script", "run_tests.bat"),
                 ("Examiner Suite", "Designed specifically for the 8-10 minute CIE-2 viva demonstration."),
                 ("Concept Proof 1", "Executes runtime polymorphism across MCQ, True/False, and Numeric questions."),
                 ("Concept Proof 2", "Compares Standard vs Negative Marking on identical test attempts."),
                 ("Concept Proof 3", "Systematically triggers and catches all 5 custom exceptions with try-catch-finally traces.")
             ], badge_text="Automated Test Harness", border_color=SUCCESS_GREEN)

    # =========================================================================
    # SLIDE 11: Individual Contribution & Work Allocation
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, LIGHT_BG)
    add_header(slide11, "Individual Work Allocation & Team Contributions")

    team_alloc = [
        ("Aditya Yadav", "Roll No: 8108",
         [("Component", "System Architecture & Base Abstractions"),
          ("Class Ownership", "abstract class Question, abstract class User, Quiz aggregation"),
          ("Key Contributions", "Designed core abstract methods (displayQuestion, checkAnswer), validated constructor parameters, and implemented User role hierarchy (Student / Instructor)."),
          ("Viva Focus", "Why use an abstract class? What is the difference between an abstract class and an interface?")]),

        ("Abhishekh Singh", "Roll No: 8104",
         [("Component", "Concrete Question Archetypes & Polymorphism"),
          ("Class Ownership", "MultipleChoiceQuestion, TrueFalseQuestion, NumericQuestion"),
          ("Key Contributions", "Implemented specialized option parsing, single-character conversion, boolean input parsing, and floating-point tolerance delta verification."),
          ("Viva Focus", "Where is runtime polymorphism used? How does dynamic method dispatch resolve question display and answer checking?")]),

        ("Priyam Raj", "Roll No: 8134",
         [("Component", "Interfaces & Pluggable Scoring Strategy"),
          ("Class Ownership", "QuizOperations, QuizEvaluator, StandardGradingPolicy, NegativeMarkingGradingPolicy"),
          ("Key Contributions", "Architected the Strategy Design Pattern allowing runtime swapping of linear vs 25% negative marking scoring models and letter grade generation."),
          ("Viva Focus", "Why did you create the QuizEvaluator interface? How does it support the Open/Closed Principle?")]),

        ("Utkarsh Chauhan", "Roll No: 8154",
         [("Component", "Exception Hierarchy & Defensive Validation"),
          ("Class Ownership", "QuizException tree (5 custom checked exceptions), InputValidator, Test Harness"),
          ("Key Contributions", "Developed domain checked exceptions, try-catch-finally handlers, scanner buffer clearing, and the automated CIE-2 test suite."),
          ("Viva Focus", "What exceptions can occur and how are they handled? What is the purpose of try, catch, and finally?")])
    ]

    coords_s11 = [
        (Inches(0.8), Inches(1.6)),
        (Inches(6.8), Inches(1.6)),
        (Inches(0.8), Inches(4.45)),
        (Inches(6.8), Inches(4.45))
    ]

    for i, (name, roll, details) in enumerate(team_alloc):
        left, top = coords_s11[i]
        add_card(slide11, left, top, Inches(5.7), Inches(2.6), f"{name} ({roll})", details, badge_text="Team Member")

    # =========================================================================
    # SLIDE 12: Future Scope, Viva Defenses & Conclusion
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, NAVY)

    # Accent bar
    bar12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar12.fill.solid()
    bar12.fill.fore_color.rgb = ROYAL_BLUE
    bar12.line.fill.background()

    # Title
    tb_c = slide12.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
    tfc = tb_c.text_frame
    pc = tfc.paragraphs[0]
    pc.text = "CIE-2 SUMMARY, FUTURE SCOPE & CONCLUSION"
    pc.font.name = "Arial"
    pc.font.size = Pt(11)
    pc.font.bold = True
    pc.font.color.rgb = RGBColor(96, 165, 250)
    pc2 = tfc.add_paragraph()
    pc2.text = "Future Project Expansion & Examiner Viva Defense"
    pc2.font.name = "Arial"
    pc2.font.size = Pt(22)
    pc2.font.bold = True
    pc2.font.color.rgb = WHITE

    # Left Card: Ongoing Project Expansion Roadmap
    add_card(slide12, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.2),
             "Ongoing Mini-Project Roadmap (Phase 2 & 3)",
             [
                 ("Phase 1 (Completed for CIE-2)", "Core Java Object-Oriented Prototype demonstrating Abstract Classes, Interfaces, Polymorphism, and Custom Exceptions."),
                 ("Phase 2: Persistent Storage", "Replace in-memory LinkedHashMap with Java Object Serialization or JDBC connectivity to MySQL / PostgreSQL."),
                 ("Phase 3: GUI & Full Mini-Project", "Migrate embedded HTTP server to Spring Boot REST APIs with JWT role-based security (Student/Faculty)."),
                 ("Advanced Assessment Features", "Per-question timer countdown using Multithreading (`ScheduledExecutorService`), question shuffling, and exportable PDF certificates.")
             ], badge_text="Expansion Roadmap", bg_color=DARK_BLUE, border_color=RGBColor(51, 65, 85))

    # Right Card: Conclusion & Viva Readiness
    add_card(slide12, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2),
             "Summary & Prepared Viva Defenses",
             [
                 ("100% Syllabus Compliance", "Direct, meaningful implementation of Unit III and Unit IV Java concepts without artificial stuffing."),
                 ("Zero-Crash Reliability", "Validated input parsing with complete recovery on illegal student choices."),
                 ("Individual Viva Readiness", "Every member is prepared with model defenses for Mrs. Trupti Najan's CIE-2 evaluation rubric."),
                 ("Demonstration Ready", "Executable via 1-click batch scripts (`run.bat`, `run_web.bat`, `run_tests.bat`)."),
                 ("Thank You!", "We welcome questions, suggestions, and feedback from the examiner.")
             ], badge_text="Viva Defense Readiness", bg_color=DARK_BLUE, border_color=ROYAL_BLUE)

    # Save to disk
    output_filename = "Online_Quiz_Management_System_CIE2.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully created: {output_filename}")

if __name__ == "__main__":
    create_deck()
